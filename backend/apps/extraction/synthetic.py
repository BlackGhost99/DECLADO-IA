"""Factures synthétiques (jeux B à E de la recette) : contenu fictif, aucun document client.

Elles servent aux tests, à la démonstration et au corpus d'évaluation synthétique. Une bonne
mesure sur ces fichiers ne prouve rien sur des factures réelles (jeu I à obtenir).
"""

import io

PAGE_W, PAGE_H = 595, 842
FONT = 9

# Colonnes du tableau (abscisse en points depuis la gauche).
COLS = {"ref": 30, "desc": 85, "code": 250, "qty": 305, "unit": 335, "price": 370, "amount": 450, "origin": 515}
HEADER = {
    "ref": "Réf.",
    "desc": "Désignation",
    "code": "HS Code",
    "qty": "Qté",
    "unit": "Unité",
    "price": "Prix unitaire",
    "amount": "Montant",
    "origin": "Origine",
}


def _esc(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_from_pages(pages: list[list[tuple[float, float, str]]], *, size: int = FONT) -> bytes:
    """PDF multipage ; chaque élément (x, top, texte) est placé depuis le coin supérieur gauche."""
    count = len(pages)
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    kids = []
    font_obj = 3 + 2 * count
    for items in pages:
        page_obj, content_obj = len(objects) + 1, len(objects) + 2
        kids.append(f"{page_obj} 0 R")
        stream = " ".join(
            f"BT /F1 {size} Tf {x:.2f} {PAGE_H - top:.2f} Td ({_esc(text)}) Tj ET" for x, top, text in items
        ).encode("cp1252", errors="replace")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {PAGE_W} {PAGE_H}] "
            f"/Resources << /Font << /F1 {font_obj} 0 R >> >> /Contents {content_obj} 0 R >>".encode()
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {count} >>".encode()
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def _row(top: float, values: dict[str, str]) -> list[tuple[float, float, str]]:
    return [(COLS[k], top, v) for k, v in values.items() if v]


def _header_block(number: str = "F-2026-001", *, supplier="Haras Fictif SARL", currency="EUR", incoterm="FOB Anvers"):
    items = [(30, 40, f"Fournisseur : {supplier}"), (30, 56, f"Facture N° {number}"), (30, 72, "Date : 15/09/2026")]
    if currency:
        items.append((30, 88, f"Devise : {currency}"))
    if incoterm:
        items.append((300, 88, f"Incoterm : {incoterm}"))
    return items


SIMPLE_LINES = [
    {"ref": "12345678", "desc": "Chevaux reproducteurs fictifs", "code": "0101.21.00", "qty": "2", "unit": "U",
     "price": "1 500,00", "amount": "3 000,00", "origin": "FR"},
    {"ref": "A-77", "desc": "Selles en cuir fictives", "code": "4201.00.00", "qty": "10", "unit": "U",
     "price": "120,50", "amount": "1 205,00", "origin": "IT"},
]  # fmt: skip


def simple_invoice_pdf(lines=None, *, total: str | None = "4 205,00", header=None, columns=None) -> bytes:
    """Jeu B : facture d'une page, code douanier présent, référence article distincte."""
    lines = SIMPLE_LINES if lines is None else lines
    cols = columns or list(HEADER)
    items = _header_block() if header is None else header
    items += _row(120, {k: HEADER[k] for k in cols})
    top = 136
    for line in lines:
        items += _row(top, {k: v for k, v in line.items() if k in cols})
        top += 14
    if total is not None:
        items += [(COLS["desc"], top + 10, "Total facture"), (COLS["amount"], top + 10, total)]
    items.append((30, 800, "Document fictif généré pour les essais."))
    return pdf_from_pages([items])


def code_cases_pdf() -> bytes:
    """Jeu C : code absent, code à six chiffres, colonne « Code » de nature incertaine."""
    header = _header_block("F-2026-002")
    items = header + _row(
        120, {"ref": "Réf.", "desc": "Désignation", "code": "Code", "qty": "Qté", "price": "P.U.", "amount": "Montant"}
    )
    rows = [
        {"ref": "87654321", "desc": "Bride en acier fictive", "code": "", "qty": "5", "price": "10,00",
         "amount": "50,00"},
        {"ref": "B-2", "desc": "Licol fictif HS Code 4201.00", "code": "", "qty": "1", "price": "8,00",
         "amount": "8,00"},
        {"ref": "C-3", "desc": "Couverture fictive", "code": "63014000", "qty": "3", "price": "20,00",
         "amount": "60,00"},
    ]  # fmt: skip
    top = 136
    for row in rows:
        items += _row(top, row)
        top += 14
    items += [(COLS["desc"], top + 10, "Total"), (COLS["amount"], top + 10, "118,00")]
    return pdf_from_pages([items])


def multipage_invoice_pdf() -> bytes:
    """Jeu D : deux pages, en-tête répété, sous-total et report, frais, remise, total."""
    page1 = _header_block("F-2026-003", incoterm="EXW Lyon")
    page1 += _row(120, HEADER)
    page1 += _row(136, {"ref": "R-1", "desc": "Brosses fictives", "code": "9603.29.00", "qty": "100", "unit": "U",
                        "price": "2,00", "amount": "200,00", "origin": "FR"})  # fmt: skip
    page1 += _row(150, {"ref": "R-2", "desc": "Étrilles fictives en plastique", "code": "9603.90.00", "qty": "50",
                        "unit": "U", "price": "3,00", "amount": "150,00", "origin": "FR"})  # fmt: skip
    page1 += _row(164, {"desc": "modèle renforcé, coloris assortis"})
    page1 += [(COLS["desc"], 190, "Sous-total page 1"), (COLS["amount"], 190, "350,00")]
    page1 += [(COLS["desc"], 204, "A reporter"), (COLS["amount"], 204, "350,00")]
    page1.append((280, 800, "Page 1/2"))
    page2 = [(30, 40, "Facture N° F-2026-003 (suite)")]
    page2 += _row(120, HEADER)
    page2 += [(COLS["desc"], 136, "Report"), (COLS["amount"], 136, "350,00")]
    page2 += _row(150, {"ref": "R-1", "desc": "Brosses fictives", "code": "9603.29.00", "qty": "100", "unit": "U",
                        "price": "2,00", "amount": "200,00", "origin": "FR"})  # fmt: skip
    page2 += [(COLS["desc"], 176, "Fret jusqu'à Anvers"), (COLS["amount"], 176, "80,00")]
    page2 += [(COLS["desc"], 190, "Remise 5 %"), (COLS["amount"], 190, "-27,50")]
    page2 += [(COLS["desc"], 204, "Total facture"), (COLS["amount"], 204, "602,50")]
    page2.append((280, 800, "Page 2/2"))
    return pdf_from_pages([page1, page2])


def unreadable_values_pdf() -> bytes:
    """Jeu T08 : désignation tronquée et quantité illisible."""
    lines = [
        {"ref": "Z-1", "desc": "Harnais d'attelage fictif double…", "code": "4201.00.00", "qty": "1O", "unit": "U",
         "price": "75,00", "amount": "750,00", "origin": "DE"},
    ]  # fmt: skip
    return simple_invoice_pdf(lines, total="750,00")


def render_png(pdf: bytes, *, dpi: int = 200, degrade: bool = False) -> bytes:
    """Rendu image d'un PDF synthétique (scan simulé) ; dégradé : flou, bruit et basse résolution."""
    import pypdfium2 as pdfium
    from PIL import ImageFilter

    doc = pdfium.PdfDocument(pdf)
    try:
        image = doc[0].render(scale=dpi / 72).to_pil().convert("L")
    finally:
        doc.close()
    if degrade:
        small = image.resize((image.width // 4, image.height // 4))
        image = small.resize(image.size).filter(ImageFilter.GaussianBlur(2.2))
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def scanned_pdf(pdf: bytes, *, dpi: int = 200, degrade: bool = False) -> bytes:
    """PDF sans couche texte (image seule), comme un scan."""
    from PIL import Image

    image = Image.open(io.BytesIO(render_png(pdf, dpi=dpi, degrade=degrade)))
    out = io.BytesIO()
    image.save(out, format="PDF", resolution=dpi)
    return out.getvalue()


def invoice_xlsx(*, external_link: bool = True, numeric_code: bool = False) -> bytes:
    """Facture Excel : codes en texte (zéros conservés), une formule avec valeur en cache, une sans."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Facture"
    ws["A1"] = "Fournisseur : Sellerie Fictive SA"
    ws["A2"], ws["B2"] = "Facture N° X-2026-9", None
    ws["A3"] = "Date : 2026-09-20"
    ws["A4"] = "Devise : USD"
    headers = ["Ref", "Description", "HS Code", "Qty", "Unit price", "Amount", "Origin"]
    for col, value in enumerate(headers, start=1):
        ws.cell(row=6, column=col, value=value)
    rows = [
        ["9001", "Fictitious saddle pads", 63079098 if numeric_code else "6307.90.98", 4, 25.5, "=D7*E7", "CN"],
        ["9002", "Fictitious bridles", "4201.00.00", 2, 40, 80, "China"],
    ]
    for r, row in enumerate(rows, start=7):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    ws["B9"], ws["F9"] = "Total", "=SUM(F7:F8)"
    if external_link:
        ws["H7"] = "='[Tarifs.xlsx]Feuil1'!A1"
    out = io.BytesIO()
    wb.save(out)
    data = out.getvalue()
    return _with_cached_values(data, {"F7": "102", "F9": "182"})


def _with_cached_values(data: bytes, values: dict[str, str]) -> bytes:
    """openpyxl n'écrit pas les résultats des formules : on simule un fichier enregistré par un tableur."""
    import re
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = body.decode()
                for ref, value in values.items():
                    text = re.sub(
                        rf'(<c r="{ref}"[^>]*>)(<f>[^<]*</f>)(<v>\s*</v>|<v/>)?',
                        lambda m, v=value: f"{m[1]}{m[2]}<v>{v}</v>",
                        text,
                    )
                body = text.encode()
            dst.writestr(info, body)
    return out.getvalue()
