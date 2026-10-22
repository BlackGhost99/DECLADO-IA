"""Factures Excel (.xlsx) : cellules lues avec leur adresse A1, formules jamais recalculées."""

import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from ..layout import Item, Page, Row, structure
from ..schema import ParsedDocument, Span, warning
from . import ExtractionFailure, Progress

MAX_SHEETS = 20
MAX_ROWS = 5000
MAX_COLS = 60


def _typed(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(repr(value))
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return str(value).strip()


def extract(path: Path, *, progress: Progress) -> ParsedDocument:
    from openpyxl import load_workbook

    data = path.read_bytes()
    try:
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionFailure(
            "DOCUMENT_UNREADABLE", "Classeur illisible ou protégé : déposer une version lisible."
        ) from exc
    pages, warnings = [], []
    external = False
    names = formulas.sheetnames[:MAX_SHEETS]
    if len(formulas.sheetnames) > MAX_SHEETS:
        warnings.append(warning("SHEETS_SKIPPED", f"Seules les {MAX_SHEETS} premières feuilles sont lues."))
    for n, name in enumerate(names):
        progress("lecture des feuilles", n, len(names))
        ws_f, ws_v = formulas[name], cached[name]
        rows = []
        for index, (row_f, row_v) in enumerate(
            zip(
                ws_f.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS),
                ws_v.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True),
                strict=False,
            )
        ):
            items = []
            for cell, value in zip(row_f, row_v, strict=False):
                raw = getattr(cell, "value", None)
                if raw is None:
                    continue
                notes = []
                if isinstance(raw, str) and raw.startswith("="):
                    if "[" in raw or "http" in raw.lower():
                        external = True
                        notes.append("Formule avec lien externe : lien non suivi.")
                    if value is None:
                        notes.append(f"Formule {raw} sans valeur enregistrée : non recalculée, valeur inconnue.")
                        text, typed = raw, None
                    else:
                        notes.append(f"Formule {raw} : valeur enregistrée par le tableur, non recalculée.")
                        typed = _typed(value)
                        text = _text(typed if typed is not None else value)
                else:
                    typed = _typed(raw)
                    text = _text(typed if typed is not None else raw)
                if not text:
                    continue
                column = cell.column
                span = Span(sheet=name, cell_range=cell.coordinate, excerpt=text[:200])
                items.append(Item(text=text, x0=column, x1=column, span=span, typed=typed, notes=notes))
            if items:
                rows.append(Row(items=items, page=None, index=index))
        pages.append(Page(number=None, rows=rows, merge_gap=0, sheet=name, exact_columns=True))
    progress("structuration", len(names), len(names))
    doc = structure(pages)
    if external:
        warnings.append(warning("EXTERNAL_LINK", "Le classeur contient des liens externes : ils n'ont pas été suivis."))
    doc.warnings.extend(warnings)
    doc.stats.update({"sheets": names})
    return doc
