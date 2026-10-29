"""Import du classeur « Tarif CEMAC 2022 SWD » (TAX-01).

Feuille 1 (une ligne par code SH à 8 chiffres et sous-position nationale) : taux DDI, DSO, DAC,
TVA, TCI, CCI, OAD, CSS, CAF, lus tels quels avec leur cellule. Feuille 2 (TEC « SH 2022 »,
droit de douane en fraction) : gardée pour comparaison, jamais appliquée. Le même import crée
l'édition du référentiel tarifaire (libellés, recherche) avec la ligne « 000 » de chaque code,
la plage réelle de la feuille étant conservée comme preuve.
"""

import io
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.db import transaction

from apps.tariffs.importers.common import IssueCollector, finish_run, register_source, source_path
from apps.tariffs.importers.csv_export import load_table
from apps.tariffs.importers.xlsx import cell_text
from apps.tariffs.models import ImportIssue, ImportRun, RawTariffRecord, SourceKind, TariffEdition

from . import rules
from .models import ScheduleLine, ScheduleRate, TaxSchedule

IMPORTER_NAME = "tarif_swd"
IMPORTER_VERSION = "1.0"
S = ImportIssue.Severity
REQUIRED = ("Code SH", "Sous position CEMAC", "Sous position nationale", *rules.SHEET_COLUMNS)
FIRST_RATE_COLUMN = 6  # G
CUMULATIVE_COLUMN = 15  # P
# Taxes dont une cellule vide signale une ligne à revoir (DAC et DSO ne frappent que certains produits).
EXPECTED = ("DDI", "TVA", "TCI", "CCI", "OAD", "CSS", "CAF")


class LayoutError(ValueError):
    pass


def _header(ws) -> list[str]:
    return [" ".join(str(c.value or "").split()) for c in ws[1]]


def is_swd_workbook(path: Path) -> bool:
    """Reconnaît la disposition SWD (utilisé par le script de démarrage pour choisir l'import)."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception:  # noqa: BLE001 - un fichier illisible n'est simplement pas un classeur SWD
        return False
    header = [" ".join(str(v or "").split()) for v in next(wb.worksheets[0].iter_rows(max_row=1, values_only=True))]
    wb.close()
    return all(h in header for h in REQUIRED)


def parse_rate(raw: str) -> Decimal | None:
    value = raw.strip().replace(",", ".")
    if not value:
        return None
    try:
        rate = Decimal(value)
    except InvalidOperation:
        return None
    return rate.normalize() if rate == rate.to_integral() else rate


def _digits(raw: str, width: int) -> str:
    text = raw.strip()
    if text.isdigit() and len(text) < width:
        text = text.zfill(width)
    return text


def _tec_reference(wb) -> tuple[str, dict[str, dict]]:
    """Feuille TEC : {code à 8 chiffres: valeurs lues}. DD est une fraction (0,05 = 5 %) ou « EXO »."""
    if len(wb.worksheets) < 2:
        return "", {}
    ws = wb.worksheets[1]
    header = _header(ws)
    if "Code SH" not in header or "DD" not in header:
        return "", {}
    col = {h: i for i, h in enumerate(header)}
    out: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2):
        code = cell_text(row[col["Code SH"]].value).replace(".", "").strip()
        sp = _digits(cell_text(row[col["SP"]].value), 2) if "SP" in col else "00"
        if len(code) != 6 or not code.isdigit() or len(sp) != 2:
            continue
        dd_raw = cell_text(row[col["DD"]].value).strip()
        if dd_raw.upper() == "EXO":
            dd = Decimal("0")
        else:
            fraction = parse_rate(dd_raw)
            dd = (fraction * rules.HUNDRED).normalize() if fraction is not None else None
        out.setdefault(
            code + sp,
            {
                "row": row[0].row,
                "dd_raw": dd_raw,
                "dd_percent": format(dd, "f") if dd is not None else None,
                "tva_raw": cell_text(row[col["TVA"]].value).strip() if "TVA" in col else "",
                "da_raw": cell_text(row[col["DA"]].value).strip() if "DA" in col else "",
            },
        )
    return ws.title, out


@transaction.atomic
def import_swd(
    path: Path,
    *,
    label: str,
    title: str,
    territory: str = "GA",
    actor=None,
    publisher_claimed: str = "",
    provenance_note: str = "",
) -> TaxSchedule:
    source = register_source(
        path,
        kind=SourceKind.XLSX,
        title=title,
        is_derived_copy=False,
        actor=actor,
        publisher_claimed=publisher_claimed,
        edition_label_claimed=label,
        provenance_note=provenance_note,
    )
    data = source_path(source).read_bytes()
    wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws_f, ws_v = wb_f.worksheets[0], wb_v.worksheets[0]
    header = _header(ws_v)
    missing = [h for h in REQUIRED if h not in header]
    if missing:
        raise LayoutError("Disposition non reconnue (colonnes absentes : " + ", ".join(missing) + ").")
    col = {h: i for i, h in enumerate(header)}

    run = ImportRun.objects.create(
        source=source, importer_name=IMPORTER_NAME, importer_version=IMPORTER_VERSION, created_by=actor
    )
    edition = TariffEdition.objects.create(label=label, territory_scope="CEMAC", import_run=run)
    issues = IssueCollector(run)
    reference_sheet, tec = _tec_reference(wb_v)

    rows = []  # (numéro de ligne, code 8, sous-position nationale, valeurs texte, formules)
    for row_f, row_v in zip(ws_f.iter_rows(min_row=2), ws_v.iter_rows(min_row=2), strict=True):
        n = row_v[0].row
        values = [cell_text(c.value) for c in row_v]
        if not any(v.strip() for v in values):
            continue
        formulas = {c.column_letter: str(c.value) for c in row_f if c.data_type == "f"}
        code = _digits(values[col["Code SH"]], 6) + _digits(values[col["Sous position CEMAC"]], 2)
        national = _digits(values[col["Sous position nationale"]], 3)
        if len(code) != 8 or not code.isdigit():
            issues.add(
                "MALFORMED_CODE",
                S.BLOCKING,
                f"Ligne {n} : code « {code} » mal formé (attendu 8 chiffres). Ignoré.",
                row=n,
            )
            continue
        rows.append((n, code, national, values, formulas))

    by_code: dict[str, list] = defaultdict(list)
    for item in rows:
        by_code[item[1]].append(item)
    primary = {}
    for code, items in by_code.items():
        primary[code] = next((i for i in items if i[2] == "000"), items[0])

    # Édition du référentiel : une ligne par code, libellé long, taux observés et cumul.
    rate_cols = [col[c] for c in rules.SHEET_COLUMNS]
    syn_header = ["Code SH", "Libellé", *rules.SHEET_COLUMNS, "Taux cumulé"]
    syn_body, syn_formulas, real_rows = [], {}, {}
    for ordinal, item in enumerate(sorted(primary.values(), key=lambda i: i[0]), start=2):
        n, code, _national, values, formulas = item
        cumulative = values[CUMULATIVE_COLUMN] if len(values) > CUMULATIVE_COLUMN else ""
        syn_body.append([code, values[col.get("Libellés longs", 3)], *(values[i] for i in rate_cols), cumulative])
        if "P" in formulas:
            syn_formulas[ordinal] = {"L": formulas["P"]}
        real_rows[ordinal] = item
    table_counts = load_table(
        run, edition, issues, syn_header, syn_body, sheet=ws_v.title, code_length=8, formulas=syn_formulas
    )
    records = list(RawTariffRecord.objects.filter(run=run, ordinal__in=real_rows))
    last = openpyxl.utils.get_column_letter(len(header))
    for record in records:
        n, _code, _national, values, formulas = real_rows[record.ordinal]
        record.cell_range, record.raw_cells, record.formulas = f"A{n}:{last}{n}", values, formulas
    header_record = RawTariffRecord.objects.get(run=run, ordinal=1)
    header_record.cell_range, header_record.raw_cells = f"A1:{last}1", header
    RawTariffRecord.objects.bulk_update([*records, header_record], ["cell_range", "raw_cells", "formulas"])

    schedule = TaxSchedule.objects.create(
        label=label,
        territory=territory,
        source=source,
        edition=edition,
        sheet=ws_v.title,
        reference_sheet=reference_sheet,
        imported_by=actor,
    )
    seen: set[tuple[str, str]] = set()
    lines, rates = [], []
    review_counts = defaultdict(int)
    for n, code, national, values, _formulas in rows:
        review = []
        if len(by_code[code]) > 1 or national != "000":
            review.append("national_subposition")
        if (code, national) in seen:
            review.append("duplicate")
            issues.add(
                "DUPLICATE_CODE",
                S.REVIEW,
                f"Code {code} sous-position {national} répété ligne {n} : la première ligne fait foi.",
                row=n,
                code=code,
            )
        seen.add((code, national))
        if any(not values[col[c]].strip() for c in EXPECTED):
            review.append("empty_tax")
        reference = tec.get(code, {})
        ddi = parse_rate(values[col["DDI"]])
        if reference and reference.get("dd_percent") is not None and ddi is not None:
            if Decimal(reference["dd_percent"]) != ddi:
                review.append("tec_difference")
        elif tec and not reference:
            review.append("not_in_reference")
        for flag in review:
            review_counts[flag] += 1
        line = ScheduleLine(
            schedule=schedule,
            hs_code=code,
            national_code=national,
            label=" ".join(values[col.get("Libellés longs", 3)].split()),
            unit=values[col["UOM_COD1"]].strip() if "UOM_COD1" in col else "",
            row_number=n,
            reference=reference,
            review=review,
        )
        lines.append(line)
        for column in rules.SHEET_COLUMNS:
            index = col[column]
            raw = values[index].strip()
            rates.append(
                ScheduleRate(
                    line=line,
                    tax=rules.COLUMN_TAX[column],
                    raw_value=raw[:40],
                    rate=parse_rate(raw),
                    cell=f"{openpyxl.utils.get_column_letter(index + 1)}{n}",
                )
            )
    ScheduleLine.objects.bulk_create(lines, batch_size=2000)
    ScheduleRate.objects.bulk_create(rates, batch_size=5000)

    unreadable = sum(1 for r in rates if r.raw_value and r.rate is None)
    if unreadable:
        issues.add("RATE_UNREADABLE", S.REVIEW, f"{unreadable} taux illisible(s) comme nombre : inconnus.")
    if review_counts["national_subposition"]:
        issues.add(
            "NATIONAL_SUBPOSITIONS",
            S.INFO,
            f"{review_counts['national_subposition']} ligne(s) de sous-positions nationales conservées dans le barème.",
        )
    if review_counts["tec_difference"]:
        issues.add(
            "TEC_DIFFERENCE",
            S.REVIEW,
            f"{review_counts['tec_difference']} ligne(s) dont le DDI diffère de la feuille « {reference_sheet} ». "
            "La feuille SWD fait foi.",
        )
    counts = {
        **table_counts,
        "sheet": ws_v.title,
        "sheets": wb_v.sheetnames,
        "columns": header,
        "schedule_lines": len(lines),
        "codes": len(by_code),
    }
    finish_run(run, counts=counts, coverage={"dimensions": ws_v.dimensions}, issues=issues)
    schedule.counts = {
        "lines": len(lines),
        "codes": len(by_code),
        "rates": len(rates),
        "empty_rates": sum(1 for r in rates if not r.raw_value),
        "reference_codes": len(tec),
        "review": dict(review_counts),
    }
    schedule.save(update_fields=["counts"])
    return schedule
