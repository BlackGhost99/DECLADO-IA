"""Dépôt, contrôle, accès temporaire et aperçu des pièces."""

import io
import logging
from datetime import datetime, timedelta
from urllib.parse import quote

from django.conf import settings
from django.core import signing
from django.utils import timezone

from apps.audit.services import record_event
from apps.core.exceptions import ApiError

from . import antivirus, quarantine
from .models import Document
from .sniff import XLSX, Readability, Rejected, sniff
from .storage import get_storage

ACCESS_SALT = "clairtransit.document-access"
logger = logging.getLogger(__name__)


def max_bytes() -> int:
    return settings.DOCUMENT_MAX_BYTES


def limits() -> dict:
    return {
        "max_bytes": max_bytes(),
        "accepted": ["PDF (.pdf)", "PNG (.png)", "JPEG (.jpg, .jpeg)", "Excel (.xlsx)"],
    }


def find_duplicates(organization, sha256: str, *, exclude_id=None):
    qs = Document.objects.for_organization(organization).filter(sha256=sha256)
    if exclude_id:
        qs = qs.exclude(pk=exclude_id)
    return list(qs.order_by("created_at")[:10])


def store_upload(
    *,
    organization,
    actor,
    upload,
    declared_mime: str = "",
    supersedes: Document | None = None,
    allow_tiff: bool = False,
) -> Document:
    """Stage -> ClamAV -> inspect -> promote; external bytes never bypass the scanner."""
    if upload.size is not None and upload.size > max_bytes():
        raise ApiError(
            "FILE_TOO_LARGE",
            f"Fichier trop volumineux ({upload.size} octets, limite {max_bytes()}).",
            413,
            {"bytes": upload.size, "max_bytes": max_bytes()},
        )
    staged = None
    stored = None
    try:
        try:
            # The bytes reach clamd while they are staged: the verdict never needs the quarantined
            # file again, which a desktop antivirus may already have locked or removed.
            with antivirus.open_stream() as scanner:
                staged = quarantine.stage(upload, max_bytes(), on_chunk=scanner.feed)
                scanner.verdict()
        except antivirus.MalwareDetected:
            logger.warning(
                "Document malveillant refusé organisation=%s sha256=%s",
                organization.id,
                getattr(staged, "sha256", None),
            )
            raise ApiError("MALWARE_DETECTED", "Fichier refusé par l'analyse antivirus.", 422) from None
        except antivirus.ScanUnavailable:
            logger.error(
                "Scanner ClamAV indisponible organisation=%s sha256=%s",
                organization.id,
                getattr(staged, "sha256", None),
            )
            raise ApiError(
                "MALWARE_SCAN_UNAVAILABLE", "Contrôle antivirus indisponible. Réessayer plus tard.", 503
            ) from None

        try:
            result = sniff(staged.path, name=upload.name, declared_mime=declared_mime, allow_tiff=allow_tiff)
        except Rejected as exc:
            status = 413 if exc.code == "FILE_TOO_LARGE" else 422
            raise ApiError(exc.code, exc.message, status, exc.details) from None

        storage = get_storage()
        stored = storage.promote_quarantined(
            staged.path, prefix=str(organization.id), sha256=staged.sha256, size=staged.size
        )
        try:
            document = Document.objects.create(
                organization=organization,
                object_key=stored.key,
                original_name=upload.name[:255],
                mime=result.mime,
                declared_mime=(declared_mime or "")[:120],
                bytes=stored.size,
                sha256=stored.sha256,
                version=supersedes.version + 1 if supersedes else 1,
                supersedes=supersedes,
                readability=result.readability,
                inspection_notes=result.notes,
                uploaded_by=actor,
            )
            record_event(
                organization=organization,
                actor=actor,
                event_type="document.stored",
                entity=document,
                current={
                    "sha256": document.sha256,
                    "bytes": document.bytes,
                    "mime": document.mime,
                    "version": document.version,
                    "readability": document.readability,
                    "supersedes": str(supersedes.id) if supersedes else None,
                    "antivirus": "clamav_clean",
                },
            )
        except Exception:
            storage.discard(stored.key)
            raise
        return document
    finally:
        if staged:
            staged.path.unlink(missing_ok=True)


def issue_access(*, document: Document, user, disposition: str) -> dict:
    """Lien temporaire, lié à l'utilisateur et à la société, émis après contrôle des droits."""
    ttl = settings.DOCUMENT_LINK_TTL_SECONDS
    token = signing.dumps(
        {"d": str(document.id), "u": str(user.id), "o": str(document.organization_id), "disp": disposition},
        salt=ACCESS_SALT,
        compress=True,
    )
    # Le nom final ne sert qu'à l'affichage (titre du lecteur, enregistrement) ; seul le jeton compte.
    name = quote(document.original_name, safe="")
    return {"url": f"/api/v1/files/{token}/{name}", "expires_at": timezone.now() + timedelta(seconds=ttl)}


def resolve_access(token: str, user) -> tuple[Document, str]:
    """Revérifie à chaque lecture : jeton valide, même utilisateur, affiliation toujours active (ACC-03)."""
    from apps.organizations.context import active_memberships

    try:
        data = signing.loads(token, salt=ACCESS_SALT, max_age=settings.DOCUMENT_LINK_TTL_SECONDS)
    except signing.BadSignature:
        raise ApiError("NOT_FOUND", "Élément introuvable.", 404) from None
    if not user.is_authenticated or str(user.id) != data.get("u"):
        raise ApiError("NOT_FOUND", "Élément introuvable.", 404)
    if not active_memberships(user).filter(organization_id=data.get("o")).exists():
        raise ApiError("NOT_FOUND", "Élément introuvable.", 404)
    document = Document.objects.filter(pk=data.get("d"), organization_id=data.get("o")).first()
    if document is None:
        raise ApiError("NOT_FOUND", "Élément introuvable.", 404)
    return document, data.get("disp", "attachment")


def _cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def xlsx_preview(document: Document, *, max_rows: int = 60, max_cols: int = 26, max_sheets: int = 10) -> dict:
    """DOC-04 : feuilles et cellules A1 en texte. Formules affichées, jamais recalculées ; liaisons non suivies."""
    if document.mime != XLSX:
        raise ApiError("PREVIEW_UNSUPPORTED", "Aperçu tabulaire réservé aux classeurs .xlsx.", 400)
    if document.readability != Readability.OK:
        raise ApiError("PREVIEW_UNAVAILABLE", "Aperçu impossible pour ce fichier.", 422)
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    # openpyxl refuse un chemin sans extension .xlsx : la clé de stockage n'en a pas.
    data = get_storage().path_for_reading(document.object_key).read_bytes()
    try:
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception:  # noqa: BLE001 - tout échec de lecture devient un état visible, pas une erreur 500
        raise ApiError("PREVIEW_UNAVAILABLE", "Le classeur n'a pas pu être lu.", 422) from None
    sheets = []
    for ws_f in formulas.worksheets[:max_sheets]:
        ws_c = cached[ws_f.title]
        rows = []
        for row_f, row_c in zip(
            ws_f.iter_rows(max_row=max_rows, max_col=max_cols),
            ws_c.iter_rows(max_row=max_rows, max_col=max_cols),
            strict=False,
        ):
            cells = []
            for cf, cc in zip(row_f, row_c, strict=False):
                if cf.value is None and cc.value is None:
                    continue
                is_formula = isinstance(cf.value, str) and cf.value.startswith("=")
                cells.append(
                    {
                        "ref": f"{get_column_letter(cf.column)}{cf.row}",
                        "value": _cell_text(cc.value),
                        "formula": cf.value if is_formula else None,
                    }
                )
            if cells:
                rows.append({"row": row_f[0].row, "cells": cells})
        sheets.append(
            {
                "name": ws_f.title,
                "dimensions": ws_f.calculate_dimension() if hasattr(ws_f, "calculate_dimension") else None,
                "rows": rows,
                "truncated": (ws_f.max_row or 0) > max_rows or (ws_f.max_column or 0) > max_cols,
            }
        )
    sheet_count = len(formulas.sheetnames)
    formulas.close()
    cached.close()
    return {
        "sheets": sheets,
        "sheet_count": sheet_count,
        "note": "Valeurs en cache lues dans le fichier ; formules affichées en texte, jamais recalculées.",
    }
