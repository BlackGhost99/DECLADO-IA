# ruff: noqa: F811  (fixtures importées de test_fiscal)
"""Exports du dossier (EXP-01 à EXP-03, T28 et T29), sur des données fictives.

Un brouillon porte un marquage visible ; seule une révision validée porte la validation interne,
jamais une acceptation par la Douane. Dans l'Excel, un code garde ses zéros et un texte « =… »
reste un texte.
"""

import io

import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.cases.models import CaseExport
from tests.test_fiscal import _private_storage, confirm, edition, schedule  # noqa: F401
from tests.test_validation import act, complete, ready_case


def export(case, fmt):
    resp = case["s"].post(f"/api/v1/cases/{case['id']}/exports", {"format": fmt})
    assert resp.status_code == 201, resp.content
    body = resp.json()
    raw = case["s"].get(f"/api/v1/cases/{case['id']}/exports/{body['id']}/file")
    assert raw.status_code == 200
    return body, b"".join(raw.streaming_content)


def pdf_text(content: bytes) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(content)
    return "\n".join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))


def test_t28_draft_then_validated_exports_carry_the_right_marking(world, schedule, edition, session):
    case, _ = ready_case(session, edition)
    detail = complete(case, world, edition)
    body, content = export(case, "pdf")
    assert body["marking"] == "draft" and body["filename"].endswith("_r1_brouillon.pdf")
    text = pdf_text(content)
    assert "BROUILLON" in text and "non validé en interne" in text
    assert "Validé en interne par" not in text
    flat = text.replace("\u00a0", " ")
    assert "54 907" in flat and "173 234" in flat  # DDI 10 % de 549 070, arrondi au franc supérieur

    rev = act(case, "submit", detail["revision"]["row_version"]).json()["revision"]
    rev = act(case, "validate", rev["row_version"], reason="Conforme (essai fictif)").json()["revision"]
    assert rev["export_count"] == 1
    body, content = export(case, "pdf")
    assert body["marking"] == "validated" and body["filename"].endswith("_r1.pdf")
    text = pdf_text(content)
    assert "Validé en interne par" in text and "pas une acceptation par la Douane" in text
    assert "BROUILLON" not in text
    listed = case["s"].get(f"/api/v1/cases/{case['id']}/exports").json()["results"]
    assert [e["marking"] for e in listed] == ["validated", "draft"]
    assert AuditEvent.objects.filter(event_type="revision.exported").count() == 2
    with pytest.raises(ValueError):
        CaseExport.objects.first().save()


def test_t29_xlsx_keeps_codes_as_text_and_neutralizes_formulas(world, schedule, edition, session):
    case, _ = ready_case(session, edition)
    complete(case, world, edition)
    rev = case["s"].get(f"/api/v1/cases/{case['id']}").json()["revision"]
    resp = case["s"].patch(f"/api/v1/revisions/{rev['id']}",
                           {"regime": '=HYPERLINK("http://x")', "row_version": rev["row_version"]})  # fmt: skip
    assert resp.status_code == 200, resp.content
    body, content = export(case, "xlsx")
    assert body["filename"].endswith("_brouillon.xlsx")
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Marchandises"]
    assert "Brouillon" in ws["A1"].value
    cells = {c.value: c for row in ws.iter_rows() for c in row if c.value is not None}
    code = cells.get("42010000") or cells.get("4201.00.00")
    assert code is not None and code.data_type == "s"
    meta = wb["Métadonnées"]
    cells = {c.value: c for row in meta.iter_rows() for c in row if c.value is not None}
    formula = cells['=HYPERLINK("http://x")']
    assert formula.data_type == "s" and formula.quotePrefix
    assert "Brouillon" in cells


def test_exports_are_bounded_to_the_company(world, schedule, edition, session):
    case, _ = ready_case(session, edition)
    complete(case, world, edition)
    body, _ = export(case, "xlsx")
    other = session("declarant@b.test")
    assert other.get(f"/api/v1/cases/{case['id']}/exports").status_code == 404
    assert other.post(f"/api/v1/cases/{case['id']}/exports", {"format": "pdf"}).status_code == 404
    assert other.get(f"/api/v1/cases/{case['id']}/exports/{body['id']}/file").status_code == 404
    assert case["s"].post(f"/api/v1/cases/{case['id']}/exports", {"format": "docx"}).status_code == 400


def test_t27_new_revision_after_validation_keeps_the_validated_one(world, schedule, edition, session):
    case, _ = ready_case(session, edition)
    detail = complete(case, world, edition)
    rev = act(case, "submit", detail["revision"]["row_version"]).json()["revision"]
    rev = act(case, "validate", rev["row_version"], reason="Conforme (essai fictif)").json()["revision"]
    export(case, "pdf")
    url = f"/api/v1/cases/{case['id']}/revisions"
    assert case["s"].post(url, {"reason": " "}).status_code == 400
    assert session("declarant@b.test").post(url, {"reason": "x"}).status_code == 404
    resp = case["s"].post(url, {"reason": "Facture rectifiée par le fournisseur (essai fictif)"})
    assert resp.status_code == 201, resp.content
    new = resp.json()["revision"]
    assert new["revision_no"] == 2 and new["state"] == "draft" and new["is_editable"]
    assert new["decisions"] == [] and new["export_count"] == 0
    assert len(new["invoices"]) == 1 and len(new["invoices"][0]["lines"]) == 1
    goods = new["goods_lines"][0]
    assert goods["decision"]["is_current"] and goods["decision"]["selected_code"] == "42010000"
    assert new["completeness"]["blocking"] == []
    # L'ancienne révision est inchangée : état, décisions, export.
    old = case["s"].get(f"/api/v1/revisions/{rev['id']}").json()
    assert old["state"] == "validated" and [d["kind"] for d in old["decisions"]] == ["submitted", "validated"]
    listed = case["s"].get(f"/api/v1/cases/{case['id']}/exports").json()["results"]
    assert [e["revision_no"] for e in listed] == [1]
    calc = case["s"].get(f"/api/v1/cases/{case['id']}/calculation").json()
    assert "frozen" not in calc and calc["articles"][0]["taxes"][0]["amount"] == "54907"
    assert AuditEvent.objects.filter(event_type="revision.created").count() == 1
    # Une révision non validée ne se double pas.
    assert case["s"].post(url, {"reason": "encore"}).status_code == 409


def test_t82_dau_project_is_prefilled_from_the_case_and_never_claims_acceptance(world, schedule, edition, session):
    """Projet au modèle DAU : cases remplies depuis le dossier, cases de la Douane vides, mention claire."""
    case, _ = ready_case(session, edition)
    complete(case, world, edition)
    body, content = export(case, "dau")
    assert body["format"] == "dau" and body["filename"].endswith("_r1_DAU_brouillon.pdf")
    assert content.startswith(b"%PDF")
    text = " ".join(pdf_text(content).replace("\u00a0", " ").split())
    assert "DÉCLARATION EN DOUANE" in text and "IMPORTATION" in text
    assert "Importateur A" in text and "Fournisseur fictif" in text  # cases 11 et 5
    assert "4201.00.00" in text  # case 37, code confirmé par l'agent
    assert "549 070" in text and "54 907" in text  # case 30 (CIF) et DDI de la case 46
    assert "ne vaut pas déclaration" in text and "BROUILLON" in text
    assert "GA/____/____/________" in text  # numéro attribué par SYDONIA, jamais inventé
    assert session("declarant@b.test").post(f"/api/v1/cases/{case['id']}/exports", {"format": "dau"}).status_code == 404
