"""Lot 2, pièces : T05 (DOC-01, DOC-02, DOC-03, DOC-05), DOC-04 et T01 sur les octets."""

import hashlib

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook

from apps.audit.models import AuditEvent
from apps.cases.models import CaseDocument
from apps.documents.models import Document
from apps.documents.synthetic import make_pdf, make_png
from apps.organizations.models import Membership

PDF = make_pdf(["FACTURE FICTIVE F-1", "Chevaux fictifs  HS Code 0101.10.00  2 x 1500,00"])


@pytest.fixture(autouse=True)
def _private_storage(settings, tmp_path):
    settings.PRIVATE_STORAGE_ROOT = tmp_path / "private"


@pytest.fixture
def case(world, session):
    s = session("declarant@a.test")
    resp = s.post("/api/v1/cases", {"title": "Pièces"})
    assert resp.status_code == 201
    return {"s": s, "id": resp.json()["id"]}


def upload(s, case_id, name, content, content_type="application/octet-stream", **data):
    headers = data.pop("headers", {})
    file = SimpleUploadedFile(name, content, content_type=content_type)
    return s.upload(f"/api/v1/cases/{case_id}/documents", {"file": file, **data}, **headers)


def stored_files(settings):
    root = settings.PRIVATE_STORAGE_ROOT / "documents"
    return [p for p in root.rglob("*") if p.is_file()] if root.exists() else []


def test_valid_pdf_is_kept_private_with_hash_author_and_type(case, settings):
    resp = upload(case["s"], case["id"], "facture.pdf", PDF, "application/pdf", usage="invoice")
    assert resp.status_code == 201, resp.content
    doc = resp.json()["document"]
    assert doc["sha256"] == hashlib.sha256(PDF).hexdigest() and doc["bytes"] == len(PDF)
    assert doc["mime"] == "application/pdf" and doc["readability"] == "ok" and doc["version"] == 1
    assert doc["uploaded_by"]["email"] == "declarant@a.test"
    [path] = stored_files(settings)
    assert path.read_bytes() == PDF and (path.stat().st_mode & 0o777) == 0o600
    assert "facture" not in path.name  # clé aléatoire, sans le nom d'origine
    assert AuditEvent.objects.filter(event_type="document.stored").exists()


@pytest.mark.parametrize(
    ("name", "content", "ctype", "status", "code"),
    [
        ("note.txt", b"bonjour", "text/plain", 422, "UNSUPPORTED_TYPE"),
        ("faux.pdf", make_png(), "application/pdf", 422, "CONTENT_TYPE_MISMATCH"),
        ("image.png", make_png(), "application/pdf", 422, "DECLARED_TYPE_MISMATCH"),
        ("vide.pdf", b"", "application/pdf", 422, "EMPTY_FILE"),
    ],
)
def test_refusals_are_explicit_and_keep_case_intact(case, settings, name, content, ctype, status, code):
    assert upload(case["s"], case["id"], "facture.pdf", PDF, "application/pdf").status_code == 201
    resp = upload(case["s"], case["id"], name, content, ctype)
    assert resp.status_code == status and resp.json()["code"] == code
    assert Document.objects.count() == 1 and len(stored_files(settings)) == 1
    assert len(case["s"].get(f"/api/v1/cases/{case['id']}/documents").json()["results"]) == 1


def test_too_large_is_refused_with_limit(case, settings):
    settings.DOCUMENT_MAX_BYTES = 100
    resp = upload(case["s"], case["id"], "facture.pdf", PDF, "application/pdf")
    assert resp.status_code == 413 and resp.json()["details"]["max_bytes"] == 100
    assert stored_files(settings) == []


@pytest.mark.parametrize(
    ("name", "content", "readability"),
    [
        ("tronque.pdf", PDF[: len(PDF) // 2], "corrupted"),
        ("protege.pdf", PDF.replace(b"/Root 1 0 R", b"/Root 1 0 R /Encrypt 9 0 R"), "protected"),
        ("chiffre.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64, "protected"),
        ("casse.xlsx", b"PK\x03\x04" + b"\0" * 64, "corrupted"),
        ("coupee.png", make_png()[:-12], "corrupted"),
    ],
)
def test_damaged_or_protected_files_are_kept_and_flagged(case, name, content, readability):
    resp = upload(case["s"], case["id"], name, content)
    assert resp.status_code == 201, resp.content
    doc = resp.json()["document"]
    assert doc["readability"] == readability and doc["previewable"] is False and doc["inspection_notes"]
    preview = case["s"].get(f"/api/v1/documents/{doc['id']}/preview").json()
    assert preview["kind"] == "unavailable"
    assert (
        case["s"].get(f"/api/v1/documents/{doc['id']}/download").status_code == 200
    )  # l'original reste téléchargeable


def test_identical_file_is_detected_then_user_decides(case, world, session):
    s = case["s"]
    first = upload(s, case["id"], "facture.pdf", PDF, usage="invoice").json()
    other_case = s.post("/api/v1/cases", {}).json()["id"]
    resp = upload(s, other_case, "copie.pdf", PDF)
    assert resp.status_code == 409 and resp.json()["code"] == "DUPLICATE_DOCUMENT"
    [existing] = resp.json()["details"]["existing"]
    assert existing["id"] == first["document"]["id"] and existing["original_name"] == "facture.pdf"
    assert Document.objects.count() == 1
    reuse = s.post(f"/api/v1/cases/{other_case}/documents/reuse", {"document_id": existing["id"], "usage": "invoice"})
    assert reuse.status_code == 201 and reuse.json()["document"]["id"] == existing["id"]
    assert upload(s, other_case, "copie.pdf", PDF, duplicate_action="keep_both").status_code == 201
    assert Document.objects.count() == 2
    # Une autre société déposant le même fichier n'apprend rien de A.
    b = session("declarant@b.test")
    b_case = b.post("/api/v1/cases", {}).json()["id"]
    assert upload(b, b_case, "x.pdf", PDF).status_code == 201
    assert b.post(f"/api/v1/cases/{b_case}/documents/reuse", {"document_id": existing["id"]}).status_code == 404


def test_idempotency_key_returns_same_attachment(case):
    headers = {"HTTP_IDEMPOTENCY_KEY": "envoi-1"}
    first = upload(case["s"], case["id"], "facture.pdf", PDF, headers=headers)
    again = upload(case["s"], case["id"], "facture.pdf", PDF, headers=headers)
    assert first.status_code == 201 and again.status_code == 200 and first.json()["id"] == again.json()["id"]
    assert CaseDocument.objects.count() == 1


def test_replacement_is_a_new_version_and_original_stays(case):
    s = case["s"]
    link = upload(s, case["id"], "facture.pdf", PDF, usage="invoice").json()
    newer = make_pdf(["FACTURE FICTIVE F-1 corrigée"])
    file = SimpleUploadedFile("facture-v2.pdf", newer, content_type="application/pdf")
    resp = s.upload(f"/api/v1/case-documents/{link['id']}/replace", {"file": file})
    assert resp.status_code == 201, resp.content
    doc = resp.json()["document"]
    assert doc["version"] == 2 and doc["supersedes_id"] == link["document"]["id"]
    links = s.get(f"/api/v1/cases/{case['id']}/documents").json()["results"]
    assert [lk["replaced_by_id"] for lk in links] == [resp.json()["id"], None]
    detail = s.get(f"/api/v1/documents/{doc['id']}").json()
    assert [v["sha256"] for v in detail["previous_versions"]] == [hashlib.sha256(PDF).hexdigest()]
    file = SimpleUploadedFile("v3.pdf", newer + b"\n", content_type="application/pdf")
    assert s.upload(f"/api/v1/case-documents/{link['id']}/replace", {"file": file}).status_code == 409


def test_temporary_link_is_bound_to_user_membership_and_time(case, world, session, settings):
    s = case["s"]
    doc_id = upload(s, case["id"], "facture.pdf", PDF).json()["document"]["id"]
    link = s.get(f"/api/v1/documents/{doc_id}/download").json()
    resp = s.client.get(link["url"])
    assert resp.status_code == 200 and b"".join(resp.streaming_content) == PDF
    assert resp["X-Content-Type-Options"] == "nosniff" and "no-store" in resp["Cache-Control"]
    assert resp["Content-Disposition"].startswith("attachment")
    # Même société, autre utilisateur : le lien ne se partage pas.
    assert session("validateur@a.test").client.get(link["url"]).status_code == 404
    # Autre société : ni métadonnées, ni lien, ni octets.
    b = session("declarant@b.test")
    assert b.get(f"/api/v1/documents/{doc_id}").status_code == 404
    assert b.get(f"/api/v1/documents/{doc_id}/download").status_code == 404
    assert b.get(f"/api/v1/documents/{doc_id}/preview").status_code == 404
    assert b.client.get(link["url"]).status_code == 404
    token, name = link["url"].removeprefix("/api/v1/files/").split("/")
    assert name == "facture.pdf"
    assert s.client.get(f"/api/v1/files/{token[:-3]}abc/{name}").status_code == 404
    # Expiration.
    settings.DOCUMENT_LINK_TTL_SECONDS = -1
    assert s.client.get(link["url"]).status_code == 404
    settings.DOCUMENT_LINK_TTL_SECONDS = 300
    # Affiliation désactivée après émission du lien : accès révoqué (ACC-03).
    Membership.objects.filter(user__email="declarant@a.test").update(active=False)
    assert s.client.get(link["url"]).status_code == 404


def test_inline_preview_for_pdf_and_image(case):
    s = case["s"]
    pdf = upload(s, case["id"], "facture.pdf", PDF).json()["document"]["id"]
    png = upload(s, case["id"], "photo.png", make_png(), "image/png").json()["document"]["id"]
    preview = s.get(f"/api/v1/documents/{pdf}/preview").json()
    assert preview["kind"] == "pdf"
    resp = s.client.get(preview["url"])
    assert resp["Content-Disposition"].startswith("inline") and resp["Content-Type"] == "application/pdf"
    assert resp["X-Frame-Options"] == "SAMEORIGIN"
    image = s.get(f"/api/v1/documents/{png}/preview").json()
    assert image["kind"] == "image"
    assert "default-src 'none'" in s.client.get(image["url"])["Content-Security-Policy"]


def test_xlsx_preview_shows_cells_and_formula_text_without_evaluating(case, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Facture"
    ws["A1"], ws["B1"], ws["C1"] = "Code", "Qté", "Total"
    ws["A2"], ws["B2"], ws["C2"] = "01011000", 2, "=B2*1500"
    ws["A3"] = '=HYPERLINK("http://exemple.invalid")'
    path = tmp_path / "facture.xlsx"
    wb.save(path)
    resp = upload(case["s"], case["id"], "facture.xlsx", path.read_bytes())
    assert resp.status_code == 201, resp.content
    preview = case["s"].get(f"/api/v1/documents/{resp.json()['document']['id']}/preview").json()
    assert preview["kind"] == "sheet"
    cells = {c["ref"]: c for r in preview["sheets"][0]["rows"] for c in r["cells"]}
    assert cells["A2"]["value"] == "01011000"  # le texte garde son zéro
    assert cells["C2"]["formula"] == "=B2*1500" and cells["C2"]["value"] is None  # aucun recalcul
    assert cells["A3"]["formula"].startswith("=HYPERLINK")
