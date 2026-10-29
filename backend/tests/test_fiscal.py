"""Lot « Taxes et règles » : TAX-01 à TAX-07, T20 à T22 et T68 à T71.

Le classeur est fictif (créé ici, taux fictifs) ; seuls la base 549 070 et le DDI 10 % viennent
d'une déclaration réelle, sans aucune donnée client, pour vérifier bases et arrondis.
"""

from decimal import Decimal

import openpyxl
import pytest

from apps.accounts.models import User
from apps.audit.models import AuditEvent
from apps.fiscal import engine, rules
from apps.fiscal.importer import import_swd, is_swd_workbook
from apps.fiscal.models import ScheduleLine, TaxAmendment
from apps.fiscal.services import activate_schedule, rates_for
from apps.tariffs.models import RawTariffRecord, TariffNode
from tests.conftest import PASSWORD
from tests.test_extraction import make_case, make_edition

HEADER = [
    "Code SH",
    "Sous position CEMAC ",
    "Sous position nationale",
    "Libellés longs",
    "Libellés courts",
    "UOM_COD1",
    *rules.SHEET_COLUMNS,
    "Taux cumulé",
]
STANDARD = ["18", "1", 0.4, 0.05, "1", 0.2]  # TVA, TCI, CCI, OAD, CSS, CAF (fictifs)


def row(code6, sub, national, label, ddi, dso=None, dac=None, rest=STANDARD):
    return [code6, sub, national, label, label, "KGM", ddi, dso, dac, *rest]


ROWS = [
    row("010121", "00", "000", "Chevaux fictifs", "5"),
    row("420100", "00", "000", "Selles fictives", "10"),
    row("260200", "00", "110", "Minerai fictif A", "10", dso="3"),
    row("260200", "00", "120", "Minerai fictif B", "20", dso="3"),
    row("220300", "00", "000", "Bière fictive", "30", dac="25"),
    row("030223", "00", "000", "Thon fictif sans TVA lue", "20", rest=[None, "1", 0.4, 0.05, "1", 0.2]),
]


def workbook(path, rows=ROWS):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TARIF CEMAC 2022 SWD"
    ws.append(HEADER)
    for n, r in enumerate(rows, start=2):
        ws.append([*r, f"=G{n}+J{n}"])
    tec = wb.create_sheet("TARIF CEMAC SH 2022")
    tec.append(["N°", "Code SH", "SP", "Désignation des produits (libellés longs)", "DD", "TVA", "DA"])
    tec.append(["01.01", None, None, "Chevaux (fictif)", None, None, None])
    tec.append([None, "0101.21", "00", "-- Reproducteurs (fictif)", 0.3, "TN", None])
    tec.append([None, "4201.00", "00", "Selles (fictif)", 0.1, "TN", None])
    wb.save(path)
    return path


@pytest.fixture(autouse=True)
def _private_storage(settings, tmp_path):
    settings.PRIVATE_STORAGE_ROOT = tmp_path / "private"


@pytest.fixture
def manager(db):
    return User.objects.create_user(email="admin-plateforme@test.test", password=PASSWORD, is_reference_manager=True)


@pytest.fixture
def schedule(db, tmp_path):
    s = import_swd(workbook(tmp_path / "swd.xlsx"), label="Barème fictif", title="swd.xlsx")
    return activate_schedule(schedule=s)


# --- T68 : import du classeur SWD ----------------------------------------------------------


def test_swd_workbook_is_imported_with_cells_review_flags_and_tariff_edition(schedule, tmp_path):
    assert is_swd_workbook(tmp_path / "swd.xlsx")
    assert schedule.is_active and schedule.counts["lines"] == 6 and schedule.counts["codes"] == 5
    horse = ScheduleLine.objects.get(schedule=schedule, hs_code="01012100")
    assert horse.row_number == 2 and horse.reference["dd_percent"] == "30" and "tec_difference" in horse.review
    ddi = horse.rates.get(tax="DDI")
    assert (ddi.raw_value, ddi.rate, ddi.cell) == ("5", Decimal("5"), "G2")
    assert "tec_difference" not in ScheduleLine.objects.get(hs_code="42010000").review
    assert "empty_tax" in ScheduleLine.objects.get(hs_code="03022300").review
    assert all("national_subposition" in ln.review for ln in ScheduleLine.objects.filter(hs_code="26020000"))
    # Édition du référentiel : une ligne par code ; la preuve garde la vraie ligne de la feuille.
    edition = schedule.edition
    assert TariffNode.objects.filter(edition=edition, kind="tariff_line").count() == 5
    record = TariffNode.objects.get(edition=edition, code="42010000").source_record
    assert record.cell_range == "A3:P3" and record.raw_cells[0] == "420100" and record.formulas == {"P": "=G3+J3"}
    assert RawTariffRecord.objects.get(run=edition.import_run, ordinal=1).raw_cells[0] == "Code SH"


def test_reimport_activates_the_new_schedule_and_keeps_the_previous(schedule, tmp_path):
    rows = [*ROWS[:1], row("420100", "00", "000", "Selles fictives", "20")]
    other = import_swd(workbook(tmp_path / "swd2.xlsx", rows), label="Barème fictif 2", title="swd2.xlsx")
    activate_schedule(schedule=other)
    schedule.refresh_from_db()
    assert not schedule.is_active and other.is_active
    assert rates_for("42010000").rates["DDI"].rate == Decimal("20")
    assert AuditEvent.objects.filter(event_type="fiscal.schedule_activated").count() == 2


def test_other_layout_is_not_a_swd_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["CODE SH", "LIBELLE", "DDI"])
    wb.save(tmp_path / "autre.xlsx")
    assert not is_swd_workbook(tmp_path / "autre.xlsx")


# --- T21 / T22 : moteur ---------------------------------------------------------------------


def test_t21_bases_follow_the_rules_and_each_tax_is_rounded_up():
    zero = dict.fromkeys(("TCI", "CCI", "OAD", "CSS", "CAF"), Decimal("0"))
    article = engine.compute_article(
        Decimal("100000"), {"DDI": Decimal("10"), "DA": Decimal("0"), "TVA": Decimal("18")} | zero
    )
    amounts = {t.code: t.amount for t in article.amounts}
    assert (amounts["DDI"], amounts["TVA"], article.total) == (Decimal("10000"), Decimal("19800"), Decimal("29800"))
    assert {t.code: t.base for t in article.amounts}["TVA"] == Decimal("110000")


def test_real_declaration_amounts_are_reproduced_to_the_franc():
    """Base 549 070 et DDI 10 % lus sur une déclaration réelle : chaque montant est celui de la déclaration."""
    rates = {"DDI": 10, "DA": 0, "TVA": 18, "TCI": 1, "CCI": "0.4", "OAD": "0.05", "CSS": 1, "CAF": "0.2"}
    article = engine.compute_article(Decimal("549070"), {k: Decimal(str(v)) for k, v in rates.items()})
    amounts = {t.code: t.amount for t in article.amounts}
    expected = {"DDI": 54907, "TVA": 108716, "TCI": 5491, "CCI": 2197, "OAD": 275, "CSS": 6040, "CAF": 1099}
    assert {k: amounts[k] for k in expected} == {k: Decimal(v) for k, v in expected.items()}
    # Total de l'article sur la déclaration (TCI à part) : 173 234.
    assert article.total - amounts["TCI"] == Decimal("173234")


def test_t22_unknown_rate_is_not_zero_and_blocks_the_total():
    rates = {"DDI": Decimal("10"), "DA": Decimal("0"), "TVA": None, "TCI": Decimal("1")}
    article = engine.compute_article(Decimal("1000"), rates)
    amounts = {t.code: t.amount for t in article.amounts}
    assert amounts["TVA"] is None and amounts["DDI"] == Decimal("100") and article.total is None
    assert "TVA" in article.unknown and article.known_total == Decimal("110")
    # Un DDI inconnu rend inconnues les taxes calculées sur valeur + DDI.
    article = engine.compute_article(Decimal("1000"), {"DDI": None, "DA": Decimal("0"), "TVA": Decimal("18")})
    tva = next(t for t in article.amounts if t.code == "TVA")
    assert tva.amount is None and tva.reason == "Dépend d'une taxe inconnue."
    # Un zéro connu donne un résultat explicite nul.
    zero = engine.compute_article(Decimal("1000"), dict.fromkeys(rules.RULES, Decimal("0")))
    assert zero.total == Decimal("0")


def test_cumulative_rate_matches_the_workbook_formula():
    rates = {"DDI": 20, "DA": 0, "TVA": 18, "TCI": 1, "CCI": "0.4", "OAD": "0.05", "CSS": 1, "CAF": "0.2"}
    assert rules.cumulative_rate({k: Decimal(str(v)) for k, v in rates.items()}) == Decimal("44.45")


# --- T20 / T69 : résolution datée et modifications ------------------------------------------


def test_resolution_reads_cells_and_never_guesses_a_national_subposition(schedule):
    beer = rates_for("22030000").rates
    assert beer["DA"].rate == Decimal("25") and beer["DA"].status == "known"
    assert rates_for("42010000").rates["DA"].status == "none"
    assert rates_for("03022300").rates["TVA"].status == "unknown"
    ore = rates_for("26020000")
    assert ore.rates["DDI"].rate is None and "Sous-position" in ore.notes[0]
    assert ore.rates["TVA"].rate == Decimal("18")  # identique sur les deux sous-positions
    assert rates_for("26020000", national="120").rates["DDI"].rate == Decimal("20")
    assert ore.rates["DSO"].status == "export_only" and ore.rates["RI"].status == "unknown"
    assert rates_for("99999999").notes == ["Code absent du barème des taxes."]


def amend(s, **data):
    body = {
        "tax": "DDI",
        "scope": "code",
        "hs_code": "4201.00.00",
        "rate": "20",
        "effective_from": "2026-01-01",
        "source_kind": "finance_law",
        "source_reference": "Loi de finances fictive 2026, art. 1",
        "reason": "Essai fictif",
    }
    return s.post("/api/v1/fiscal/amendments", body | data)


def test_t69_only_the_platform_admin_amends_a_rate_and_it_is_dated_and_audited(world, schedule, manager, session):
    resp = amend(session("admin@a.test"))
    assert resp.status_code == 403
    s = session(manager.email)
    assert amend(s, reason="").status_code == 400
    assert amend(s, source_reference="").json()["details"]["field"] == "source_reference"
    assert amend(s, rate="120").status_code == 400
    assert amend(s, hs_code="1234").status_code == 400
    assert amend(s, hs_code="9999.99.99").json()["code"] == "UNKNOWN_CODE"
    resp = amend(s)
    assert resp.status_code == 201, resp.content
    body = resp.json()
    assert (body["sequence"], body["rate"], body["previous_rate"], body["version"]) == (1, "20", "10", 2)
    event = AuditEvent.objects.get(event_type="fiscal.rate_amended")
    assert event.organization is None and event.current["rate"] == "20" and event.previous == {"rate": "10"}
    # La modification vaut à partir de sa date ; avant, la cellule du barème reste appliquée.
    from datetime import date

    assert rates_for("42010000", on=date(2025, 12, 31)).rates["DDI"].rate == Decimal("10")
    after = rates_for("42010000", on=date(2026, 6, 1))
    assert after.rates["DDI"].rate == Decimal("20") and after.rates["DDI"].source == "amendment" and after.version == 2
    # Le barème importé n'est pas réécrit et une modification ne se modifie pas.
    assert ScheduleLine.objects.get(hs_code="42010000").rates.get(tax="DDI").rate == Decimal("10")
    with pytest.raises(ValueError):
        TaxAmendment.objects.get().save()


def test_redevance_is_set_for_all_codes_and_a_later_amendment_wins(world, schedule, manager, session):
    s = session(manager.email)
    assert amend(s, tax="RI", hs_code="", rate="0.5").status_code == 201
    assert amend(s, tax="TVA", scope="all", hs_code="", rate="19").status_code == 201
    assert amend(s, tax="TVA", rate="5").status_code == 201
    rates = rates_for("42010000").rates
    assert rates["RI"].rate == Decimal("0.5") and rates["TVA"].rate == Decimal("5")
    assert rates_for("01012100").rates["TVA"].rate == Decimal("19")
    overview = s.get("/api/v1/fiscal/taxes").json()
    assert overview["version"] == 4 and overview["can_edit"] is True
    ri = next(t for t in overview["taxes"] if t["code"] == "RI")
    assert ri["status"] == "set" and ri["common_rate"] == "0.5"
    assert amend(s, tax="DSO").status_code == 400


def test_overview_review_and_code_endpoints_are_readable_but_not_editable_by_companies(world, schedule, session):
    s = session("declarant@a.test")
    overview = s.get("/api/v1/fiscal/taxes").json()
    assert overview["can_edit"] is False and overview["schedule"]["label"] == "Barème fictif"
    ddi = next(t for t in overview["taxes"] if t["code"] == "DDI")
    assert ddi["base_kind"] == "value" and ddi["distinct"] == 4
    tva = next(t for t in overview["taxes"] if t["code"] == "TVA")
    assert tva["base_kind"] == "value_ddi_da" and tva["empty"] == 1
    review = s.get("/api/v1/fiscal/taxes/review/tec_difference").json()
    assert review["count"] == 1 and review["results"][0]["hs_code"] == "01012100"
    assert s.get("/api/v1/fiscal/taxes/review/autre").status_code == 404
    detail = s.get("/api/v1/fiscal/rates", code="0101.21.00").json()
    assert detail["label"] == "Chevaux fictifs" and detail["row_number"] == 2
    assert [r["tax"] for r in detail["rates"]][:3] == ["DDI", "DA", "TVA"]
    assert detail["cumulative_rate"] == "26.6"  # 5 + 105 × 18 % + 1 + 0,4 + 0,05 + 105 × 1 % + 0,2
    assert s.get("/api/v1/fiscal/rates", code="0101").status_code == 400
    assert s.post("/api/v1/fiscal/amendments", {"tax": "DDI"}).status_code == 403


# --- T70 / T71 : calcul d'un dossier --------------------------------------------------------


@pytest.fixture
def edition(db):
    return make_edition()


def goods_with_code(case, *, amount, currency="XAF", incoterm="CIF", code="4201.00.00"):
    s = case["s"]
    inv = s.post(
        f"/api/v1/revisions/{case['revision']}/invoices",
        {"supplier_name": "Fournisseur fictif", "invoice_number": "F-1", "currency": currency, "incoterm": incoterm},
    ).json()
    line = s.post(
        f"/api/v1/invoices/{inv['id']}/lines",
        {"description_raw": "Selles fictives", "quantity": "1", "unit_price": amount, "amount": amount}
        | ({"customs_code": code} if code else {}),
    )
    assert line.status_code == 201, line.content
    goods = s.post(f"/api/v1/revisions/{case['revision']}/goods-lines", {"invoice_line_ids": [line.json()["id"]]})
    assert goods.status_code == 201, goods.content
    return goods.json()


def confirm(case, goods, edition):
    s = case["s"]
    check = s.post(f"/api/v1/goods-lines/{goods['id']}/code-checks", {"edition_id": str(edition.id)}).json()
    resp = s.post(
        f"/api/v1/goods-lines/{goods['id']}/decisions",
        {
            "action": "confirm",
            "code_check_id": check["id"],
            "reason": "Code de la facture conforme (essai fictif).",
            "goods_line_version": s.get(f"/api/v1/goods-lines/{goods['id']}").json()["row_version"],
        },
    )
    assert resp.status_code == 201, resp.content


def test_t70_case_calculation_shows_each_tax_with_base_rate_and_cell(world, schedule, edition, session):
    case = make_case(session, edition)
    goods = goods_with_code(case, amount="549070")
    url = f"/api/v1/cases/{case['id']}/calculation"
    before = case["s"].get(url).json()
    [article] = before["articles"]
    assert article["code_status"] == "missing" and article["total"] is None
    assert article["taxes"][0]["reason"] == "Code SH non confirmé."
    confirm(case, goods, edition)
    body = case["s"].get(url).json()
    [article] = body["articles"]
    assert article["hs_code"] == "42010000" and article["value"]["xaf"] == "549070"
    taxes = {t["code"]: t for t in article["taxes"]}
    assert (taxes["DDI"]["amount"], taxes["DDI"]["rate"], taxes["DDI"]["cell"]) == ("54907", "10", "G3")
    assert (taxes["TVA"]["base"], taxes["TVA"]["amount"]) == ("603977", "108716")
    assert taxes["DA"]["rate_status"] == "none" and taxes["DA"]["amount"] == "0"
    assert article["total"] == "178725" and body["provisional"] is False
    # Redevance informatique sans taux saisi : total de la déclaration absent, somme connue affichée.
    assert body["totals"]["total"] is None and body["totals"]["known_total"] == "178725"
    assert body["global_taxes"][0]["code"] == "RI" and body["global_taxes"][0]["amount"] is None
    assert body["rules_version"] == rules.RULES_VERSION and body["taxes_version"] == 1


def test_t70_fob_invoice_without_freight_piece_is_provisional_and_euro_uses_the_fixed_parity(
    world, schedule, edition, session
):
    case = make_case(session, edition)
    goods = goods_with_code(case, amount="1000", currency="EUR", incoterm="FOB Anvers")
    confirm(case, goods, edition)
    body = case["s"].get(f"/api/v1/cases/{case['id']}/calculation").json()
    assert body["articles"][0]["value"]["xaf"] == "655957"
    assert body["provisional"] is True
    assert any("fret absente" in r for r in body["provisional_reasons"])
    assert any("assurance absente" in r for r in body["provisional_reasons"])


def test_unknown_currency_gives_an_unknown_value_not_zero(world, schedule, edition, session):
    case = make_case(session, edition)
    goods = goods_with_code(case, amount="1000", currency="USD")
    confirm(case, goods, edition)
    [article] = case["s"].get(f"/api/v1/cases/{case['id']}/calculation").json()["articles"]
    assert article["value"]["xaf"] is None and article["total"] is None
    assert "taux de change à saisir" in article["value"]["notes"][0]


def test_t71_other_company_cannot_read_the_calculation(world, schedule, edition, session):
    case = make_case(session, edition)
    goods_with_code(case, amount="1000")
    assert session("declarant@b.test").get(f"/api/v1/cases/{case['id']}/calculation").status_code == 404
    assert session().get(f"/api/v1/cases/{case['id']}/calculation").status_code in (401, 403)
