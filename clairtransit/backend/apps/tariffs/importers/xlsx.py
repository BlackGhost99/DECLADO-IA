"""Import d'un classeur tarifaire Excel original (REF-01, REF-02).

Le classeur est lu deux fois avec openpyxl : formules d'une part, valeurs en cache de
l'autre. Aucune formule n'est évaluée et les liaisons externes ne sont jamais ouvertes :
leur cible est seulement relevée. Les valeurs issues d'une formule liée à un fichier
absent restent des observations signalées comme telles.
"""

import io
import re
from collections import Counter
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.db import transaction

from ..models import ImportIssue, ImportRun, SourceKind, TariffEdition
from .common import IssueCollector, finish_run, register_source, source_path
from .csv_export import load_table

IMPORTER_NAME = "xlsx"
IMPORTER_VERSION = "1.0"
S = ImportIssue.Severity
EXTERNAL_REF = re.compile(r"\[(\d+)\]")
# LOOKUP(clé, …) : la clé devrait être la cellule de code de la même ligne.
LOOKUP_KEY = re.compile(r"^=LOOKUP\(([^,]+),", re.IGNORECASE)


def cell_text(value) -> str:
    """Valeur en cache rendue en texte sans perte : zéros initiaux et décimales conservés."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "VRAI" if value else "FAUX"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(Decimal(repr(value)).normalize(), "f")
    return str(value)


@transaction.atomic
def import_xlsx(
    path: Path,
    *,
    edition_label: str,
    title: str,
    sheet: str | None = None,
    code_length: int = 8,
    actor=None,
    publisher_claimed: str = "",
    provenance_note: str = "",
    territory: str = "CEMAC",
) -> TariffEdition:
    source = register_source(
        path,
        kind=SourceKind.XLSX,
        title=title,
        is_derived_copy=False,
        actor=actor,
        publisher_claimed=publisher_claimed,
        edition_label_claimed=edition_label,
        provenance_note=provenance_note,
    )
    run = ImportRun.objects.create(
        source=source, importer_name=IMPORTER_NAME, importer_version=IMPORTER_VERSION, created_by=actor
    )
    edition = TariffEdition.objects.create(label=edition_label, territory_scope=territory, import_run=run)
    issues = IssueCollector(run)

    stored = source_path(source)
    wb_formulas = openpyxl.load_workbook(
        io.BytesIO(stored.read_bytes()), data_only=False, read_only=False, keep_links=True
    )
    wb_values = openpyxl.load_workbook(
        io.BytesIO(stored.read_bytes()), data_only=True, read_only=False, keep_links=True
    )
    ws = wb_formulas[sheet] if sheet else wb_formulas.worksheets[0]
    wv = wb_values[ws.title]

    links = []
    for index, link in enumerate(wb_formulas._external_links, start=1):
        target = getattr(getattr(link, "file_link", None), "Target", "") or ""
        links.append({"index": index, "target": target})
    if links:
        issues.add(
            "EXTERNAL_LINK",
            S.REVIEW,
            f"Le classeur référence {len(links)} fichier(s) externe(s) absent(s) : jamais ouvert(s) ni recalculé(s). "
            "Les valeurs concernées sont les caches enregistrés dans le classeur.",
            links=links,
        )

    header: list[str] = []
    body: list[list[str]] = []
    formulas: dict[int, dict[str, str]] = {}
    external_by_column: Counter[str] = Counter()
    misaligned: Counter[str] = Counter()
    misaligned_examples: dict[str, list[dict]] = {}
    text_numbers: Counter[str] = Counter()
    examples: dict[str, str] = {}

    for row_f, row_v in zip(ws.iter_rows(), wv.iter_rows(), strict=True):
        n = row_f[0].row
        values = []
        row_formulas = {}
        for cf, cv in zip(row_f, row_v, strict=True):
            if cf.data_type == "f":
                formula = str(cf.value)
                row_formulas[cf.column_letter] = formula
                if EXTERNAL_REF.search(formula):
                    external_by_column[cf.column_letter] += 1
                    examples.setdefault(cf.column_letter, f"{cf.coordinate} {formula}")
                key = LOOKUP_KEY.match(formula)
                if key and key.group(1).strip().replace("$", "") != f"A{n}":
                    misaligned[cf.column_letter] += 1
                    column_examples = misaligned_examples.setdefault(cf.column_letter, [])
                    if len(column_examples) < 5:
                        column_examples.append({"cell": cf.coordinate, "formula": formula})
            elif n > 1 and isinstance(cv.value, str) and cv.value.strip() and cf.column > 2:
                try:
                    Decimal(cv.value.strip())
                    text_numbers[cf.column_letter] += 1
                except ArithmeticError:
                    pass
            values.append(cell_text(cv.value))
        if row_formulas:
            formulas[n] = row_formulas
        if n == 1:
            header = values
        else:
            body.append(values)

    table_counts = load_table(
        run, edition, issues, header, body, sheet=ws.title, code_length=code_length, formulas=formulas
    )

    letters = {openpyxl.utils.get_column_letter(i + 1): h for i, h in enumerate(header)}
    for column, count in sorted(external_by_column.items()):
        issues.add(
            "EXTERNAL_FORMULA",
            S.REVIEW,
            f"Colonne « {letters.get(column, column)} » ({column}) : {count} valeur(s) issue(s) d'une formule liée "
            "à un fichier absent. Valeur en cache conservée, non vérifiable.",
            column=letters.get(column, column),
            letter=column,
            count=count,
            example=examples[column],
        )
    for column, count in sorted(misaligned.items()):
        issues.add(
            "LOOKUP_KEY_MISMATCH",
            S.REVIEW,
            f"Colonne « {letters.get(column, column)} » ({column}) : {count} formule(s) RECHERCHE dont la clé n'est "
            "pas le code de la même ligne. La valeur peut correspondre à un autre article.",
            column=letters.get(column, column),
            letter=column,
            count=count,
            examples=misaligned_examples[column],
        )
    for column, count in sorted(text_numbers.items()):
        issues.add(
            "NUMBER_STORED_AS_TEXT",
            S.INFO,
            f"Colonne « {letters.get(column, column)} » ({column}) : {count} nombre(s) saisi(s) comme texte.",
            column=letters.get(column, column),
            letter=column,
            count=count,
        )

    counts = {
        **table_counts,
        "sheet": ws.title,
        "sheets": wb_formulas.sheetnames,
        "columns": header,
        "formula_cells": sum(len(f) for f in formulas.values()),
        "external_links": links,
    }
    coverage = {"rows_without_node": len(body) - table_counts["tariff_lines"], "dimensions": ws.dimensions}
    finish_run(run, counts=counts, coverage=coverage, issues=issues)
    return edition
